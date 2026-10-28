"""Table structure: header detection, widths, alignment, nested tables."""

from __future__ import annotations

import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_helpers import (
    document_xml,
    full_text,
    load,
    paragraph_with,
    table_rows_text,
)
from scripts.markdown_to_word import MAX_IMAGE_WIDTH_INCHES, markdown_to_word


def test_gfm_table_structure_and_header(convert):
    source = (
        "| 列A | 列B |\n"
        "| --- | --- |\n"
        "| 1 | 2 |\n"
        "| 3 | 4 |\n"
    )
    result = convert(source)

    assert result.tables == 1
    assert table_rows_text(load(result.output_path)) == [
        ["列A", "列B"],
        ["1", "2"],
        ["3", "4"],
    ]
    xml = document_xml(result.output_path)
    assert xml.count("<w:tblHeader") == 1

    document = load(result.output_path)
    header_cell = document.tables[0].rows[0].cells[0]
    assert all(run.bold for run in header_cell.paragraphs[0].runs)


def test_table_without_header_row_is_not_styled_as_header(convert):
    source = (
        "<table>\n"
        "<tr><td>a</td><td>b</td></tr>\n"
        "<tr><td>c</td><td>d</td></tr>\n"
        "</table>\n"
    )
    result = convert(source)

    xml = document_xml(result.output_path)
    assert "<w:tblHeader" not in xml

    document = load(result.output_path)
    first_cell = document.tables[0].rows[0].cells[0]
    assert not any(run.bold for run in first_cell.paragraphs[0].runs)


def test_table_data_rows_alternate_shading(convert):
    source = (
        "| h |\n"
        "| --- |\n"
        "| r1 |\n"
        "| r2 |\n"
        "| r3 |\n"
    )
    result = convert(source)

    document = load(result.output_path)
    table = document.tables[0]
    shading = []
    for row in table.rows[1:]:
        cell_properties = row.cells[0]._tc.tcPr
        element = cell_properties.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd"
        )
        shading.append(
            element.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill"
            )
            if element is not None
            else None
        )
    assert shading[0] is None
    assert shading[1] == "F7F9FB"
    assert shading[2] is None


def test_nested_table_is_rendered_inside_the_cell(convert):
    source = (
        "<table>\n"
        "<tr><td>外层</td><td><table><tr><td>内层</td></tr></table></td></tr>\n"
        "</table>\n"
    )
    result = convert(source)

    document = load(result.output_path)
    assert len(document.tables) == 1
    outer = document.tables[0]
    assert len(outer.rows) == 1
    nested = outer.rows[0].cells[1].tables
    assert len(nested) == 1
    assert nested[0].rows[0].cells[0].text == "内层"
    assert result.tables == 2
    parent_width = outer.rows[0].cells[1].width.inches
    nested_width = sum(cell.width.inches for cell in nested[0].rows[0].cells)
    assert nested_width <= parent_width


def test_html_table_cell_whitespace_is_normalized(convert):
    source = (
        "<table>\n"
        "<tr><th>列</th></tr>\n"
        "<tr>\n"
        "  <td>\n    多个   空格\n  </td>\n"
        "</tr>\n"
        "</table>\n"
    )
    result = convert(source)

    document = load(result.output_path)
    text = document.tables[0].rows[1].cells[0].text
    assert text == "多个 空格"
    assert "\n" not in text


def test_column_widths_fill_page_width(convert):
    source = (
        "| a | b | c |\n"
        "| --- | --- | --- |\n"
        "| 1 | 2 | 3 |\n"
    )
    result = convert(source)

    document = load(result.output_path)
    widths = [
        document.tables[0].rows[0].cells[index].width.inches
        for index in range(3)
    ]
    assert sum(widths) == pytest.approx(MAX_IMAGE_WIDTH_INCHES, abs=0.02)
    assert all(width >= 0.4 for width in widths)


def test_table_cell_alignment(convert):
    source = (
        "| 左 | 中 | 右 |\n"
        "| :--- | :---: | ---: |\n"
        "| 1 | 2 | 3 |\n"
    )
    result = convert(source)

    document = load(result.output_path)
    row = document.tables[0].rows[1]
    assert row.cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert row.cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert row.cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_table_cell_inline_formatting(convert):
    source = (
        "| 名称 | 说明 |\n"
        "| --- | --- |\n"
        "| `code` | **粗体** |\n"
    )
    result = convert(source)

    document = load(result.output_path)
    code_cell = document.tables[0].rows[1].cells[0]
    run = code_cell.paragraphs[0].runs[0]
    assert run.font.name == "Consolas"

    bold_cell = document.tables[0].rows[1].cells[1]
    assert bold_cell.paragraphs[0].runs[0].bold is True


def test_table_followed_by_text_keeps_order(convert):
    source = (
        "| a |\n"
        "| --- |\n"
        "| 1 |\n"
        "\n"
        "表格后文本\n"
    )
    result = convert(source)

    text = full_text(result.output_path)
    assert "表格后文本" in text
    assert paragraph_with(result.output_path, "表格后文本").style.name == "Normal"


def test_rows_with_fewer_cells_are_padded(convert):
    source = (
        "<table>\n"
        "<tr><td>a</td><td>b</td></tr>\n"
        "<tr><td>c</td></tr>\n"
        "</table>\n"
    )
    result = convert(source)

    document = load(result.output_path)
    assert len(document.tables) == 1
    rows = table_rows_text(document)
    assert rows[0] == ["a", "b"]
    assert rows[1][0] == "c"
    assert len(rows[1]) == 2


def test_multi_line_cell_keeps_breaks(convert):
    source = "| 列 |\n| --- |\n| 第一行<br>第二行 |\n"
    result = convert(source)

    document = load(result.output_path)
    cell_text = document.tables[0].rows[1].cells[0].text
    assert "第一行" in cell_text and "第二行" in cell_text
    assert "\n" in cell_text


def test_empty_table_is_skipped(convert):
    result = convert("<table></table>\n")

    assert result.tables == 0


def test_single_column_table(convert):
    source = "| 只有一列 |\n| --- |\n| 值 |\n"
    result = convert(source)
    document = load(result.output_path)

    assert len(document.tables[0].columns) == 1
    width = document.tables[0].rows[0].cells[0].width.inches
    assert width == pytest.approx(MAX_IMAGE_WIDTH_INCHES, abs=0.02)


def test_three_level_nested_tables(convert):
    source = (
        "<table><tr><td>层1"
        "<table><tr><td>层2"
        "<table><tr><td>层3</td></tr></table>"
        "</td></tr></table>"
        "</td></tr></table>\n"
    )
    result = convert(source)
    document = load(result.output_path)

    assert result.tables == 3
    level2 = document.tables[0].rows[0].cells[0].tables
    assert len(level2) == 1
    level3 = level2[0].rows[0].cells[0].tables
    assert len(level3) == 1
    assert level3[0].rows[0].cells[0].text == "层3"


def test_nested_table_font_is_smaller(convert):
    source = (
        "<table><tr><td>x"
        "<table><tr><td>inner</td></tr></table>"
        "</td></tr></table>\n"
    )
    result = convert(source)
    document = load(result.output_path)
    outer_cell = document.tables[0].rows[0].cells[0]
    inner_cell = outer_cell.tables[0].rows[0].cells[0]

    outer_size = outer_cell.paragraphs[0].runs[0].font.size.pt
    inner_size = inner_cell.paragraphs[0].runs[0].font.size.pt
    assert inner_size < outer_size


def test_large_table_structure(convert):
    header = "| " + " | ".join(f"h{index}" for index in range(8)) + " |"
    divider = "| " + " | ".join("---" for _ in range(8)) + " |"
    body = [
        "| " + " | ".join(f"{row}-{col}" for col in range(8)) + " |"
        for row in range(10)
    ]
    source = "\n".join([header, divider, *body]) + "\n"
    result = convert(source)

    document = load(result.output_path)
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 11
    assert len(document.tables[0].columns) == 8
    assert "9-7" in full_text(result.output_path)


def test_table_in_blockquote(convert):
    source = "> | a |\n> | --- |\n> | 1 |\n"
    result = convert(source)

    assert result.tables == 1
    assert "1" in full_text(result.output_path)


def test_two_tables_in_a_row(convert):
    source = (
        "| a |\n| --- |\n| 1 |\n\n"
        "| b |\n| --- |\n| 2 |\n"
    )
    result = convert(source)

    assert result.tables == 2
    document = load(result.output_path)
    assert len(document.tables) == 2


def test_table_cell_with_only_code(convert):
    source = "| 代码 |\n| --- |\n| `x = 1` |\n"
    result = convert(source)
    document = load(result.output_path)
    run = document.tables[0].rows[1].cells[0].paragraphs[0].runs[0]

    assert run.text == "x = 1"
    assert run.font.name == "Consolas"
