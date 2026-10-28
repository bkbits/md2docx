"""Unit-level tests for internal branches that end-to-end flows rarely hit."""

from __future__ import annotations

import os
import subprocess
from pathlib import Path

import pytest
from bs4 import BeautifulSoup

import scripts.markdown_to_word as converter_module
from docx_helpers import full_text, load, media_names, texts
from scripts.markdown_to_word import (
    MAX_IMAGE_WIDTH_INCHES,
    FormulaRenderer,
    MarkdownDocxConverter,
    MermaidRenderer,
    calculate_column_widths,
    markdown_to_word,
    pygments_style,
    tag_classes,
)


def _renderer(tmp_path, mode: str = "auto", timeout: int = 30) -> MermaidRenderer:
    return MermaidRenderer(
        tmp_path, color="2980B9", font="SimSun", mode=mode, timeout_seconds=timeout
    )


@pytest.mark.parametrize(
    ("available", "expected_prefix"),
    [
        ({"mmdc": "/bin/mmdc", "bunx": "/bin/bunx", "npx": "/bin/npx"}, ["/bin/mmdc"]),
        ({"bunx": "/bin/bunx", "npx": "/bin/npx"}, ["/bin/bunx", "--bun"]),
        ({"npx": "/bin/npx"}, ["/bin/npx", "--yes"]),
        ({}, None),
    ],
)
def test_find_command_priority(monkeypatch, available, expected_prefix):
    monkeypatch.setattr(
        converter_module.shutil, "which", lambda name: available.get(name)
    )

    command = MermaidRenderer._find_command()

    if expected_prefix is None:
        assert command is None
    else:
        assert command is not None
        assert command[: len(expected_prefix)] == expected_prefix


@pytest.mark.skipif(os.name != "nt", reason="Windows-only Chrome lookup")
def test_find_chrome_windows_install_path(monkeypatch):
    monkeypatch.setattr(converter_module.shutil, "which", lambda name: None)
    monkeypatch.setattr(converter_module.Path, "is_file", lambda self: True)

    chrome = MermaidRenderer._find_chrome()

    assert chrome is not None
    assert chrome.lower().endswith((".exe",))


def test_find_chrome_missing(monkeypatch):
    monkeypatch.setattr(converter_module.shutil, "which", lambda name: None)
    monkeypatch.setattr(converter_module.Path, "is_file", lambda self: False)

    assert MermaidRenderer._find_chrome() is None


def test_mermaid_timeout_returns_warning(monkeypatch, tmp_path):
    monkeypatch.setattr(
        MermaidRenderer, "_find_command", staticmethod(lambda: ["mmdc"])
    )
    renderer = _renderer(tmp_path, mode="auto", timeout=3)
    monkeypatch.setattr(
        converter_module.subprocess,
        "run",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            subprocess.TimeoutExpired("mmdc", 3)
        ),
    )

    path, warning = renderer.render("graph TD\n A --> B", 1)

    assert path is None
    assert "timed out after 3 seconds" in warning


def test_mermaid_timeout_required_raises(monkeypatch, tmp_path):
    monkeypatch.setattr(
        MermaidRenderer, "_find_command", staticmethod(lambda: ["mmdc"])
    )
    renderer = _renderer(tmp_path, mode="required", timeout=3)
    monkeypatch.setattr(
        converter_module.subprocess,
        "run",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            subprocess.TimeoutExpired("mmdc", 3)
        ),
    )

    with pytest.raises(RuntimeError):
        renderer.render("graph TD\n A --> B", 1)


def test_mermaid_batch_timeout_falls_back_to_per_diagram(monkeypatch, tmp_path):
    monkeypatch.setattr(
        MermaidRenderer, "_find_command", staticmethod(lambda: ["mmdc"])
    )
    renderer = _renderer(tmp_path, mode="auto", timeout=3)
    monkeypatch.setattr(
        converter_module.subprocess,
        "run",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            subprocess.TimeoutExpired("mmdc", 3)
        ),
    )

    results = renderer.render_all(["a", "b"])

    assert len(results) == 2
    assert all(path is None for path, _ in results)
    assert all("timed out" in warning for _, warning in results)


def test_mermaid_batch_missing_output_falls_back(monkeypatch, tmp_path):
    monkeypatch.setattr(
        MermaidRenderer, "_find_command", staticmethod(lambda: ["mmdc"])
    )
    renderer = _renderer(tmp_path, mode="auto", timeout=3)

    class Completed:
        returncode = 0
        stderr = ""
        stdout = ""

    monkeypatch.setattr(
        converter_module.subprocess, "run", lambda *args, **kwargs: Completed()
    )

    results = renderer.render_all(["a", "b"])

    assert all(path is None for path, _ in results)
    assert all("failed" in warning for _, warning in results)


def test_table_cell_with_nested_table_and_text(tmp_path):
    source = (
        "<table><tr><td>\n"
        "  前文\n"
        "  <table><tr><td>内层</td></tr></table>\n"
        "  后文\n"
        "</td></tr></table>\n"
    )
    result = markdown_to_word(
        markdown_content=source,
        output_path=tmp_path / "mixed.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    document = load(result.output_path)
    cell = document.tables[0].rows[0].cells[0]
    assert [paragraph.text for paragraph in cell.paragraphs if paragraph.text] == [
        "前文",
        "后文",
    ]
    assert len(cell.tables) == 1
    assert "内层" in full_text(result.output_path)


def test_html_section_block_is_recursed(tmp_path):
    source = "<section><p>章节正文</p></section>\n"
    result = markdown_to_word(
        markdown_content=source,
        output_path=tmp_path / "section.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert "章节正文" in full_text(result.output_path)


def test_unknown_block_tag_falls_back_to_text(tmp_path):
    source = "<figure>图注内容</figure>\n"
    result = markdown_to_word(
        markdown_content=source,
        output_path=tmp_path / "figure.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert "图注内容" in full_text(result.output_path)


def test_ordered_list_with_invalid_start(tmp_path):
    source = '<ol start="abc"><li>项目</li></ol>\n'
    result = markdown_to_word(
        markdown_content=source,
        output_path=tmp_path / "start.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert result.lists == 1
    assert "项目" in full_text(result.output_path)


def test_list_style_missing_falls_back_to_manual_marker(tmp_path, monkeypatch):
    from docx.document import Document as DocxDocument

    original_add = DocxDocument.add_paragraph

    def patched(self, text="", style=None):
        if style is not None:
            raise KeyError(style)
        return original_add(self, text, style)

    monkeypatch.setattr(DocxDocument, "add_paragraph", patched)
    result = markdown_to_word(
        markdown_content="1. 第一项\n2. 第二项\n",
        output_path=tmp_path / "marker.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    lines = [line for line in texts(result.output_path) if line.strip()]
    assert lines == ["1. 第一项", "2. 第二项"]


def test_heading_with_softbreak_gets_combined_id(tmp_path):
    source = "标题一\n标题二\n=====\n\n[跳转](#标题一-标题二)\n"
    result = markdown_to_word(
        markdown_content=source,
        output_path=tmp_path / "setext.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    from docx_helpers import anchors, bookmarks

    assert set(anchors(result.output_path)) <= set(bookmarks(result.output_path))
    assert anchors(result.output_path)


def test_column_widths_enforce_minimum(monkeypatch):
    rows = BeautifulSoup(
        "<table><tr>"
        + "<th>窄</th>"
        + "".join("<th>很宽的一列内容</th>" for _ in range(9))
        + "</tr></table>",
        "html.parser",
    ).find_all("tr")
    widths = calculate_column_widths(rows, 10)

    assert sum(widths) == pytest.approx(MAX_IMAGE_WIDTH_INCHES, abs=0.01)
    assert min(widths) >= 0.4
    assert all(width > 0 for width in widths)


def test_hyperlink_preserves_surrounding_spaces(tmp_path):
    result = markdown_to_word(
        markdown_content="[ 带空格 ](https://example.com)\n",
        output_path=tmp_path / "spaces.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    from docx_helpers import document_xml

    assert 'xml:space="preserve"' in document_xml(result.output_path)


def test_bold_link_keeps_bold_formatting(tmp_path):
    result = markdown_to_word(
        markdown_content="**[粗链接](https://example.com)**\n",
        output_path=tmp_path / "bold-link.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    from docx_helpers import document_xml

    xml = document_xml(result.output_path)
    hyperlink = xml[xml.index("<w:hyperlink") :]
    assert "<w:b/>" in hyperlink or '<w:b w:val="1"/>' in hyperlink


def test_pygments_style_without_pygments(monkeypatch):
    monkeypatch.setattr(converter_module, "Token", None)

    assert pygments_style(object()) == {}


def test_inline_math_with_escaped_dollar(tmp_path):
    result = markdown_to_word(
        markdown_content="公式 $a\\$b$ 结束\n",
        output_path=tmp_path / "escaped.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert result.formulas_rendered + result.formulas_fallback == 1
    assert "结束" in full_text(result.output_path)


def test_block_formula_falls_back_to_shaded_text(tmp_path, monkeypatch):
    monkeypatch.setattr(FormulaRenderer, "render", lambda self, latex, block: None)
    result = markdown_to_word(
        markdown_content="$$\na^2\n$$\n",
        output_path=tmp_path / "formula-fallback.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert result.formulas_rendered == 0
    assert result.formulas_fallback == 1
    paragraph = next(
        item for item in load(result.output_path).paragraphs if item.text.strip()
    )
    assert "a^2" in paragraph.text
    assert paragraph.alignment is not None


def test_svg_is_converted_when_converter_succeeds(tmp_path, monkeypatch, png_bytes):
    from types import SimpleNamespace

    svg = tmp_path / "v.svg"
    svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>',
        encoding="utf-8",
    )

    def fake_run(command, **kwargs):
        output = Path(command[command.index("-o") + 1])
        output.write_bytes(png_bytes)
        return SimpleNamespace(returncode=0, stderr="", stdout="")

    monkeypatch.setattr(converter_module.shutil, "which", lambda name: "/bin/rsvg-convert")
    monkeypatch.setattr(converter_module.subprocess, "run", fake_run)
    result = markdown_to_word(
        markdown_content=f"![矢量]({svg.as_posix()})\n",
        output_path=tmp_path / "svg-ok.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert result.images == 1
    assert media_names(result.output_path)


def test_svg_conversion_failure_warns(tmp_path, monkeypatch):
    from types import SimpleNamespace

    svg = tmp_path / "v.svg"
    svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>',
        encoding="utf-8",
    )
    monkeypatch.setattr(converter_module.shutil, "which", lambda name: "/bin/rsvg-convert")
    monkeypatch.setattr(
        converter_module.subprocess,
        "run",
        lambda command, **kwargs: SimpleNamespace(
            returncode=1, stderr="rsvg boom", stdout=""
        ),
    )
    result = markdown_to_word(
        markdown_content=f"![矢量]({svg.as_posix()})\n",
        output_path=tmp_path / "svg-fail.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert result.images == 0
    assert any("rsvg boom" in warning for warning in result.warnings)


def test_tag_classes_with_string_attribute():
    class Stub:
        def get(self, key, default=None):
            return "alpha beta" if key == "class" else default

    assert tag_classes(Stub()) == ["alpha", "beta"]

    class Empty:
        def get(self, key, default=None):
            return default

    assert tag_classes(Empty()) == []


def test_media_parts_registered_for_images(tmp_path):
    result = markdown_to_word(
        markdown_content=(
            "![一](data:image/png;base64,"
            "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAEklEQVR4nGM8YWPDwMDAxAAGAA+2AUSPIZanAAAAAElFTkSuQmCC)\n"
        ),
        output_path=tmp_path / "media.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert result.images == 1
    assert len(media_names(result.output_path)) == 1


def test_italic_link_keeps_italic(tmp_path):
    result = markdown_to_word(
        markdown_content="*[斜链接](https://example.com)*\n",
        output_path=tmp_path / "italic-link.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    from docx_helpers import document_xml

    xml = document_xml(result.output_path)
    hyperlink = xml[xml.index("<w:hyperlink") :]
    assert "<w:i/>" in hyperlink or '<w:i w:val="1"/>' in hyperlink


@pytest.mark.parametrize(
    ("token", "expected_color"),
    [
        ("Keyword", "0000CC"),
        ("String", "A31515"),
        ("Comment", "008000"),
        ("Number", "098658"),
        ("Operator", "666666"),
        ("Error", "B91C1C"),
    ],
)
def test_pygments_token_colors(token, expected_color):
    from pygments import token as tokens

    token_type = getattr(tokens.Token, token)

    assert pygments_style(token_type).get("color") == expected_color


def test_pygments_unknown_token_uses_default_color():
    from pygments.token import Text

    assert pygments_style(Text) == {"color": "1F2937"}


def test_blockquote_with_bare_text(tmp_path):
    source = "<blockquote>裸文本引用</blockquote>\n"
    result = markdown_to_word(
        markdown_content=source,
        output_path=tmp_path / "quote.docx",
        overwrite=True,
        mermaid_mode="skip",
    )

    assert "裸文本引用" in full_text(result.output_path)


def test_cli_runs_as_script(tmp_path):
    """Exercise the ``__main__`` entry point in a real process."""
    import subprocess as sp
    import sys

    root = Path(__file__).resolve().parents[1]
    output = tmp_path / "script.docx"
    completed = sp.run(
        [
            sys.executable,
            str(root / "scripts" / "markdown_to_word.py"),
            "--markdown-content",
            "# 标题",
            "--output",
            str(output),
            "--mermaid",
            "skip",
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )

    assert completed.returncode == 0, completed.stderr
    assert "Created:" in completed.stdout
    assert output.exists()
