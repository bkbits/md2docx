"""Helpers for asserting on generated DOCX files."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def load(path: str | Path):
    return Document(str(path))


def paragraphs(path: str | Path) -> list[Paragraph]:
    return list(load(path).paragraphs)


def texts(path: str | Path) -> list[str]:
    """Body-level paragraph texts, in document order."""
    return [paragraph.text for paragraph in load(path).paragraphs]


def non_empty_texts(path: str | Path) -> list[str]:
    return [text for text in texts(path) if text.strip()]


def full_text(path: str | Path) -> str:
    """Every paragraph text in the body, including nested table cells."""
    document = load(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        chunks.extend(_table_texts(table))
    return "\n".join(chunks)


def _table_texts(table) -> list[str]:
    chunks: list[str] = []
    for row in table.rows:
        for cell in row.cells:
            chunks.extend(paragraph.text for paragraph in cell.paragraphs)
            for nested in cell.tables:
                chunks.extend(_table_texts(nested))
    return chunks


def paragraph_with(path: str | Path, needle: str) -> Paragraph:
    for paragraph in load(path).paragraphs:
        if needle in paragraph.text:
            return paragraph
    raise AssertionError(f"no paragraph contains {needle!r}")


def runs_with(path: str | Path, needle: str):
    for paragraph in load(path).paragraphs:
        for run in paragraph.runs:
            if needle in run.text:
                yield run
    return


def single_run(path: str | Path, needle: str):
    runs = list(runs_with(path, needle))
    assert runs, f"no run contains {needle!r}"
    return runs[0]


def document_xml(path: str | Path) -> str:
    with zipfile.ZipFile(str(path)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def settings_xml(path: str | Path) -> str:
    with zipfile.ZipFile(str(path)) as archive:
        return archive.read("word/settings.xml").decode("utf-8")


def numbering_xml(path: str | Path) -> str:
    with zipfile.ZipFile(str(path)) as archive:
        return archive.read("word/numbering.xml").decode("utf-8")


def media_names(path: str | Path) -> list[str]:
    with zipfile.ZipFile(str(path)) as archive:
        return [
            name
            for name in archive.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        ]


def bookmarks(path: str | Path) -> dict[str, str]:
    """Map bookmark name -> w:id for every bookmark in the document."""
    document = load(path)
    result: dict[str, str] = {}
    for start in document.element.body.iter(qn("w:bookmarkStart")):
        name = start.get(qn("w:name"))
        if name:
            result[name] = start.get(qn("w:id"))
    return result


def anchors(path: str | Path) -> list[str]:
    document = load(path)
    return [
        element.get(qn("w:anchor")) or ""
        for element in document.element.body.iter(qn("w:hyperlink"))
    ]


def bordered_paragraphs(path: str | Path) -> list[tuple[str, list[str]]]:
    """(text, edge names) for each paragraph carrying borders."""
    result = []
    for paragraph in load(path).paragraphs:
        properties = paragraph._p.pPr
        if properties is None:
            continue
        borders = properties.find(qn("w:pBdr"))
        if borders is None:
            continue
        edges = [child.tag.split("}")[1] for child in borders]
        result.append((paragraph.text, edges))
    return result


def paragraph_alignment(path: str | Path, needle: str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = paragraph_with(path, needle)
    assert paragraph.alignment is not None
    return WD_ALIGN_PARAGRAPH(paragraph.alignment)


def style_fonts(document, style_name: str) -> dict[str, str | None]:
    style = document.styles[style_name]
    rfonts = style._element.rPr.rFonts
    return {
        "ascii": rfonts.get(qn("w:ascii")) if rfonts is not None else None,
        "eastAsia": rfonts.get(qn("w:eastAsia")) if rfonts is not None else None,
    }


def table_rows_text(document, index: int = 0) -> list[list[str]]:
    table = document.tables[index]
    return [
        [cell.text for cell in row.cells]
        for row in table.rows
    ]


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def _read(archive: zipfile.ZipFile, name: str) -> str:
    return archive.read(name).decode("utf-8")


def relationships(path: str | Path) -> dict[str, str]:
    """Map relationship id -> target for the main document part."""
    import xml.etree.ElementTree as ElementTree

    with zipfile.ZipFile(str(path)) as archive:
        root = ElementTree.fromstring(_read(archive, "word/_rels/document.xml.rels"))
    return {
        element.get("Id"): element.get("Target")
        for element in root
        if element.get("Id")
    }


def style_ids(path: str | Path) -> set[str]:
    import xml.etree.ElementTree as ElementTree

    with zipfile.ZipFile(str(path)) as archive:
        root = ElementTree.fromstring(_read(archive, "word/styles.xml"))
    return {
        element.get(f"{{{W_NS}}}styleId")
        for element in root.iter(f"{{{W_NS}}}style")
    }


def used_style_ids(path: str | Path) -> list[str]:
    document = load(path)
    return [
        element.get(qn("w:val")) or ""
        for element in document.element.body.iter(qn("w:pStyle"))
    ]


def used_numbering_ids(path: str | Path) -> set[str]:
    document = load(path)
    return {
        element.get(qn("w:val")) or ""
        for element in document.element.body.iter(qn("w:numId"))
    }


def defined_numbering_ids(path: str | Path) -> set[str]:
    document = load(path)
    numbering = document.part.numbering_part.element
    return {
        element.get(qn("w:numId")) or ""
        for element in numbering.iter(qn("w:num"))
    }


def drawing_count(path: str | Path) -> int:
    return len(document_xml(path).split("<w:drawing")) - 1


def image_descriptions(path: str | Path) -> list[str]:
    document = load(path)
    return [
        element.get("descr") or ""
        for element in document.element.body.iter(f"{{{WP_NS}}}docPr")
    ]


def image_extents(path: str | Path) -> list[tuple[int, int]]:
    """(cx, cy) in EMU for every inline picture."""
    document = load(path)
    return [
        (int(element.get("cx")), int(element.get("cy")))
        for element in document.element.body.iter(f"{{{WP_NS}}}extent")
    ]


def heading_paragraphs(path: str | Path) -> list[tuple[int, str]]:
    result = []
    for paragraph in load(path).paragraphs:
        name = paragraph.style.name
        if name.startswith("Heading "):
            result.append((int(name.split()[1]), paragraph.text))
    return result


def link_targets(path: str | Path) -> list[str]:
    """Targets of external hyperlinks, resolved through the relationship part."""
    relations = relationships(path)
    document = load(path)
    targets = []
    for hyperlink in document.element.body.iter(qn("w:hyperlink")):
        relationship_id = hyperlink.get(qn("r:id"))
        if relationship_id:
            targets.append(relations.get(relationship_id, ""))
    return targets


def assert_document_integrity(path: str | Path) -> None:
    """Check OOXML references so broken documents fail loudly in tests."""
    import xml.etree.ElementTree as ElementTree

    with zipfile.ZipFile(str(path)) as archive:
        assert archive.testzip() is None
        for name in (
            "word/document.xml",
            "word/styles.xml",
            "word/settings.xml",
            "word/numbering.xml",
            "word/_rels/document.xml.rels",
        ):
            ElementTree.fromstring(_read(archive, name))

    known_styles = style_ids(path)
    for style_id in used_style_ids(path):
        assert style_id in known_styles, f"unknown pStyle {style_id!r}"

    defined_numbers = defined_numbering_ids(path)
    for number_id in used_numbering_ids(path):
        assert number_id in defined_numbers, f"unknown numId {number_id!r}"

    relations = relationships(path)
    document = load(path)
    for hyperlink in document.element.body.iter(qn("w:hyperlink")):
        relationship_id = hyperlink.get(qn("r:id"))
        if relationship_id:
            assert relationship_id in relations, f"missing rel {relationship_id!r}"
    for blip in document.element.body.iter(f"{{{A_NS}}}blip"):
        embed_id = blip.get(f"{{{R_NS}}}embed")
        if embed_id:
            assert embed_id in relations, f"missing image rel {embed_id!r}"

    starts = {
        element.get(qn("w:id"))
        for element in document.element.body.iter(qn("w:bookmarkStart"))
    }
    ends = {
        element.get(qn("w:id"))
        for element in document.element.body.iter(qn("w:bookmarkEnd"))
    }
    assert starts == ends, f"bookmark mismatch {starts ^ ends}"
    assert len(starts) == len(
        [
            element
            for element in document.element.body.iter(qn("w:bookmarkStart"))
        ]
    )

    doc_pr_ids = [
        element.get("id")
        for element in document.element.body.iter(f"{{{WP_NS}}}docPr")
    ]
    assert len(doc_pr_ids) == len(set(doc_pr_ids)), (
        f"duplicate drawing ids {doc_pr_ids}"
    )
